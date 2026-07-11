from collections.abc import Callable
from io import BytesIO

from openpyxl.workbook import Workbook as WorkbookType

from services.owui import download_file
from subservers.xlsx import _ops

SCRIPT_API = """
Run a Python script that builds or edits the current project. The script is
sandboxed: no imports, no file or network access. Available functions:

- write_rows(sheet: str, rows: list[list], anchor: str = "A1",
             style: str | None = None) -> None
  Fill a contiguous block of cells from row-major `rows`, starting at the
  `anchor` A1 reference. Numbers stay numeric, a string starting with `=`
  becomes a formula, and None clears the cell. An optional named cell style
  from `styles` (see the `create_project` / `open_project` result) formats
  every written cell.
- write_cell(sheet: str, ref: str, value=None,
             style: str | None = None) -> None
  Write one cell by A1 reference, optionally with a named style.
- await add_image(sheet: str, file_id: str, anchor: str = "A1") -> None
  Insert an image the user attached in OpenWebUI (async: must be
  awaited) with its top-left corner on the `anchor` cell. `file_id` is
  never invented.
- add_chart(sheet: str, kind: str, data: str, categories: str,
            title: str | None = None, anchor: str = "A1") -> None
  Insert a native Excel chart with its top-left corner on the `anchor`
  cell. `kind` is "bar", "line" or "pie". `data` is the A1 range holding
  the values, one column per series, with the series name in the first
  row (e.g. "B1:C4"); `categories` is the range of the category labels
  (e.g. "A2:A4"). The chart stays linked to those cells, so it updates
  when they change. A pie chart takes exactly one series column.
- add_comment(sheet: str, ref: str, text: str) -> None
  Attach a review comment (a cell note) to one cell, replacing any
  existing note there. The cell's value stays untouched — use it to give
  feedback on an opened workbook.
- read_sheet(sheet: str) -> list[list[str]]
  The sheet's used range as text rows (formulas as their formula string).
- list_sheets() -> list[dict]
  Worksheets in order (title, used rows/cols).
- add_sheet(title: str, index: int | None = None) -> None
- move_sheet(title: str, to_index: int) -> None
- remove_sheets(titles: list[str]) -> None

Plain Python works: loops, conditionals, f-strings, comprehensions,
functions, print. `print(...)` output and the last expression are returned
as `output`. If the script fails, fix it and call `run_script` again.

Example:

    data = [["Region", "Umsatz"], ["Nord", 1.2], ["Sued", 0.9]]
    write_rows("Sheet", data)
    write_cell("Sheet", f"B{len(data) + 1}", f"=SUM(B2:B{len(data)})")
    add_chart("Sheet", "bar", data="B1:B3", categories="A2:A3", anchor="D2")

Column widths are auto-fitted on `finalize_project`. After the requested
batch of edits, call `finalize_project` once (not after each change).
""".strip()


def script_functions(
    workbook: WorkbookType,
    token: str,
) -> dict[str, Callable]:

    def write_rows(
        sheet: str,
        rows: list[list[bool | int | float | str | None]],
        anchor: str = "A1",
        style: str | None = None,
    ) -> None:
        _ops.write_rows(workbook, sheet, rows, anchor, style)

    def write_cell(
        sheet: str,
        ref: str,
        value: bool | int | float | str | None = None,
        style: str | None = None,
    ) -> None:
        _ops.write_cell(workbook, sheet, ref, value, style)

    async def add_image(
        sheet: str,
        file_id: str,
        anchor: str = "A1",
    ) -> None:

        file_content = await download_file(file_id=file_id, token=token)

        try:
            _ops.insert_picture(workbook, sheet, BytesIO(file_content), anchor)
        except OSError as error:
            raise ValueError(
                f"File '{file_id}' is not a supported image."
            ) from error

    def add_chart(
        sheet: str,
        kind: str,
        data: str,
        categories: str,
        title: str | None = None,
        anchor: str = "A1",
    ) -> None:
        _ops.add_chart(workbook, sheet, kind, data, categories, title, anchor)

    def add_comment(sheet: str, ref: str, text: str) -> None:
        _ops.add_comment(workbook, sheet, ref, text)

    def read_sheet(sheet: str) -> list[list[str]]:
        return _ops.read_sheet(workbook, sheet)

    def list_sheets() -> list[dict]:
        return [
            info.model_dump()
            for info in _ops.list_sheet_infos(workbook)
        ]

    def add_sheet(title: str, index: int | None = None) -> None:
        _ops.add_sheet(workbook, title, index)

    def move_sheet(title: str, to_index: int) -> None:
        _ops.move_sheet(workbook, title, to_index)

    def remove_sheets(titles: list[str]) -> None:
        _ops.remove_sheets(workbook, titles)

    return {
        "write_rows": write_rows,
        "write_cell": write_cell,
        "add_image": add_image,
        "add_chart": add_chart,
        "add_comment": add_comment,
        "read_sheet": read_sheet,
        "list_sheets": list_sheets,
        "add_sheet": add_sheet,
        "move_sheet": move_sheet,
        "remove_sheets": remove_sheets,
    }
